"""PowerPoint（.pptx）解析器：按「页 → 要点」切分。

每页标题作为一级标题（形成独立章节），正文按要点切分；页内表格按行语义化
转写。惰性导入 ``python-pptx``。
"""

from __future__ import annotations

from pathlib import Path

from app.rag.ingest.parser.base import ParserDependencyError, table_rows_to_blocks
from app.rag.models import Block, BlockKind, DocType, StructuredDoc


def parse_pptx(path: Path, doc_id: str, *, tenant_id: str = "", kb_id: str = "") -> StructuredDoc:
    """解析 PPT 为结构化文档（页码写入溯源元数据）。"""
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise ParserDependencyError("解析 PPT 需要 python-pptx，请执行 pip install python-pptx") from exc

    presentation = Presentation(str(path))
    blocks: list[Block] = []

    for slide_no, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        for shape in slide.shapes:
            if shape.has_table:
                rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                rows = [row for row in rows if any(cell for cell in row)]
                if rows:
                    blocks.extend(
                        table_rows_to_blocks(rows[0], rows[1:], slide_no=slide_no, start_row=2)
                    )
                continue
            if not shape.has_text_frame:
                continue
            # PPT 正文以要点为主，统一按 BULLET 处理，利于切分后保留要点边界
            is_title = title_shape is not None and shape.shape_id == title_shape.shape_id
            for paragraph in shape.text_frame.paragraphs:
                text = "".join(run.text for run in paragraph.runs).strip() or paragraph.text.strip()
                if not text:
                    continue
                blocks.append(
                    Block(
                        kind=BlockKind.HEADING if is_title else BlockKind.BULLET,
                        text=text,
                        level=1 if is_title else 0,
                        slide_no=slide_no,
                    )
                )

    title = next((b.text for b in blocks if b.kind is BlockKind.HEADING), path.stem)
    return StructuredDoc(
        doc_id=doc_id,
        doc_type=DocType.PPTX,
        title=title,
        blocks=blocks,
        tenant_id=tenant_id,
        kb_id=kb_id,
    )
