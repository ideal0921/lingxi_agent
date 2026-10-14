"""Word（.docx）解析器：按 Heading 1~6 标题树递归切分，表格按行语义化转写。

惰性导入 ``python-docx``；通过遍历 ``body`` 子元素保证段落与表格的**原始顺序**
（``doc.paragraphs`` 与 ``doc.tables`` 分列会丢失文档交错顺序）。
"""

from __future__ import annotations

import re
from pathlib import Path

from app.rag.ingest.parser.base import ParserDependencyError, table_rows_to_blocks
from app.rag.models import Block, BlockKind, DocType, StructuredDoc

# 兼容中英文 Word 的内置样式名：Heading 1 / 标题 1
_HEADING_STYLE_PATTERN = re.compile(r"(?:Heading|标题)\s*(\d)")
_LIST_STYLE_KEYWORDS = ("list", "列表")


def parse_docx(path: Path, doc_id: str, *, tenant_id: str = "", kb_id: str = "") -> StructuredDoc:
    """解析 Word 文档为结构化文档。"""
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:
        raise ParserDependencyError("解析 Word 需要 python-docx，请执行 pip install python-docx") from exc

    document = Document(str(path))
    blocks: list[Block] = []

    for child in document.element.body.iterchildren():
        # 标签形如 {namespace}p / {namespace}tbl，按局部名分派
        local_name = child.tag.rsplit("}", 1)[-1]
        if local_name == "p":
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            level = _heading_level(paragraph.style.name if paragraph.style else "")
            if level:
                blocks.append(Block(kind=BlockKind.HEADING, text=text, level=level))
            else:
                style_name = (paragraph.style.name if paragraph.style else "").lower()
                kind = BlockKind.BULLET if any(key in style_name for key in _LIST_STYLE_KEYWORDS) else BlockKind.PARAGRAPH
                blocks.append(Block(kind=kind, text=text))
        elif local_name == "tbl":
            table = Table(child, document)
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            rows = [row for row in rows if any(cell for cell in row)]
            if rows:
                blocks.extend(table_rows_to_blocks(rows[0], rows[1:], start_row=1))

    title = next((b.text for b in blocks if b.kind is BlockKind.HEADING), path.stem)
    return StructuredDoc(
        doc_id=doc_id,
        doc_type=DocType.DOCX,
        title=title,
        blocks=blocks,
        tenant_id=tenant_id,
        kb_id=kb_id,
    )


def _heading_level(style_name: str) -> int:
    """由 Word 样式名解析标题层级（Heading 1~6），非标题返回 0。"""
    matched = _HEADING_STYLE_PATTERN.search(style_name)
    if not matched:
        return 0
    return min(int(matched.group(1)), 6)
