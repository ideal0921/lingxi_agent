"""Excel / CSV 解析器：按「表头 → 行」语义化转写，禁止按单元格切分。

- CSV：标准库 ``csv`` 解析，零依赖；
- Excel：惰性导入 ``openpyxl``，缺失时抛 :class:`ParserDependencyError`；
  每个工作表生成一级标题块，使切分后的 ``heading_path`` 可定位到具体表。
"""

from __future__ import annotations

import csv
from pathlib import Path

from app.rag.ingest.parser.base import (
    ParserDependencyError,
    table_rows_to_blocks,
)
from app.rag.models import Block, BlockKind, DocType, StructuredDoc


def parse_csv(path: Path, doc_id: str, *, tenant_id: str = "", kb_id: str = "") -> StructuredDoc:
    """解析 CSV：首行作表头，其余行语义化转写。"""
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = [row for row in csv.reader(handle) if any(cell.strip() for cell in row)]
    header = rows[0] if rows else []
    blocks = table_rows_to_blocks(header, rows[1:], sheet=path.stem, start_row=2)
    return StructuredDoc(
        doc_id=doc_id,
        doc_type=DocType.SHEET,
        title=path.stem,
        blocks=blocks,
        tenant_id=tenant_id,
        kb_id=kb_id,
    )


def parse_xlsx(path: Path, doc_id: str, *, tenant_id: str = "", kb_id: str = "") -> StructuredDoc:
    """解析 Excel：逐工作表转写，工作表名作为一级标题。"""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # 惰性导入：避免整个入库管道被重依赖拖垮
        raise ParserDependencyError("解析 Excel 需要 openpyxl，请执行 pip install openpyxl") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    blocks: list[Block] = []
    try:
        for worksheet in workbook.worksheets:
            rows = [
                ["" if cell is None else str(cell) for cell in row]
                for row in worksheet.iter_rows(values_only=True)
            ]
            rows = [row for row in rows if any(cell.strip() for cell in row)]
            if not rows:
                continue
            # 工作表名作为一级标题，保证 heading_path 可区分多表
            blocks.append(Block(kind=BlockKind.HEADING, text=worksheet.title, level=1))
            blocks.extend(table_rows_to_blocks(rows[0], rows[1:], sheet=worksheet.title, start_row=2))
    finally:
        workbook.close()

    return StructuredDoc(
        doc_id=doc_id,
        doc_type=DocType.SHEET,
        title=path.stem,
        blocks=blocks,
        tenant_id=tenant_id,
        kb_id=kb_id,
    )
